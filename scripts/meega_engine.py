"""Extrai o MEEGA+ diretamente do Excel e calcula as tabelas por item.

Uso: python extrair_meega.py --entrada "respostas.xlsx"
Dependência: openpyxl. A planilha é aberta somente para leitura.
"""
import argparse
from collections import Counter
import csv
import json
from pathlib import Path
import statistics
import unicodedata

import openpyxl

BASE = Path(__file__).resolve().parent
ENTRADA = BASE.parent / "TotalMeltdown_ Questionário de Avaliação (JSS) (respostas).xlsx"


def normalizar(texto):
    return " ".join(unicodedata.normalize("NFKC", str(texto)).split()).casefold()


def localizar_colunas(cabecalhos, nomes):
    """Usa os textos dos cabeçalhos, independentemente da posição das colunas."""
    resultado = {}
    for nome in nomes:
        encontrados = [i for i, valor in enumerate(cabecalhos)
                       if valor is not None and normalizar(valor) == normalizar(nome)]
        if len(encontrados) != 1:
            raise ValueError(f"Cabeçalho ausente ou duplicado: {nome!r}")
        resultado[nome] = encontrados[0]
    return resultado


def ler_planilha(entrada, mapa=BASE / "itens_meega.json"):
    config = json.loads(Path(mapa).read_text(encoding="utf-8"))
    itens = config["items"]
    if len({i["item"] for i in itens}) != len(itens):
        raise ValueError("Identificadores de itens duplicados no mapeamento.")
    arquivo = openpyxl.load_workbook(entrada, read_only=True, data_only=True)
    respostas, participantes = [], set()
    try:
        aba = arquivo[config["sheet"]]
        linhas = aba.iter_rows(values_only=True)
        colunas = localizar_colunas(next(linhas),
            [config["id_header"], config["consent_header"]] + [i["header"] for i in itens])
        for numero, linha in enumerate(linhas, start=2):
            if all(v is None or str(v).strip() == "" for v in linha):
                continue
            identificador = linha[colunas[config["id_header"]]]
            if identificador is None or not str(identificador).strip():
                raise ValueError(f"ID ausente na linha {numero}.")
            if isinstance(identificador, (int, float)) and not isinstance(identificador, bool):
                identificador = str(int(identificador)) if int(identificador) == identificador else str(identificador)
            identificador = str(identificador).strip()
            if identificador in participantes:
                raise ValueError(f"ID duplicado: {identificador}.")
            consentimento = linha[colunas[config["consent_header"]]]
            if normalizar(consentimento) != normalizar(config["accepted_consent"]):
                raise ValueError(f"Consentimento não confirmado na linha {numero}.")
            participantes.add(identificador)
            for item in itens:
                indice = colunas[item["header"]]
                valor = linha[indice]
                celula = f"{openpyxl.utils.get_column_letter(indice + 1)}{numero}"
                if isinstance(valor, bool) or not isinstance(valor, (int, float)) or valor not in (1, 2, 3, 4, 5):
                    raise ValueError(f"Resposta ausente ou fora de 1--5 em {celula}: {valor!r}")
                respostas.append({"participant_id": identificador, "item": item["item"],
                                  "rating": int(valor), "source_cell": celula})
    finally:
        arquivo.close()
    if not respostas:
        raise ValueError("Nenhuma resposta encontrada.")
    return itens, respostas


def resumir(itens, respostas):
    resumo = []
    for item in itens:
        valores = [r["rating"] for r in respostas if r["item"] == item["item"]]
        n = len(valores)
        if not n:
            raise ValueError(f"Item sem respostas: {item['item']}")
        contagem = Counter(valores)
        r = {**item, "n": n, "mean": statistics.mean(valores),
             "median": statistics.median(valores),
             "sd": statistics.stdev(valores) if n > 1 else ""}
        r.update({f"n_{k}": contagem[k] for k in range(1, 6)})
        r.update({f"pct_{k}": 100 * contagem[k] / n for k in range(1, 6)})
        for nome, categorias in [("1_2", (1, 2)), ("3", (3,)), ("4_5", (4, 5))]:
            total = sum(contagem[k] for k in categorias)
            r[f"ratings_{nome}_n"] = total
            r[f"ratings_{nome}_pct"] = 100 * total / n
        resumo.append(r)
    return resumo


def salvar_tabelas(saida, itens, respostas, resumo):
    saida = Path(saida)
    saida.mkdir(parents=True, exist_ok=True)
    for nome, registros in [("meega_respostas.csv", respostas), ("meega_itens.csv", resumo)]:
        with (saida / nome).open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=list(registros[0]), delimiter=";")
            writer.writeheader()
            writer.writerows(registros)

    def tex(texto):
        return str(texto).replace("&", r"\&").replace("%", r"\%").replace("_", r"\_")

    n = len({r["participant_id"] for r in respostas})
    linhas = [r"\begin{table*}[htbp]", r"\centering", r"\small",
        f"\\caption{{Item-level responses to the adapted MEEGA+ questionnaire ($n={n}$). Ratings range from 1 to 5; the last three columns report percentages. Labels are abbreviated.}}",
        r"\label{tab:meega-items}", r"\begin{tabular}{llrrrrr}", r"\toprule",
        r"Item & Statement & Mean & Median & 1--2 (\%) & 3 (\%) & 4--5 (\%) \\", r"\midrule"]
    painel_anterior = None
    for r in resumo:
        if r["panel"] != painel_anterior:
            if painel_anterior is not None:
                linhas.append(r"\addlinespace")
            linhas.append(r"\multicolumn{7}{l}{\textit{" + tex(r["panel"]) + r"}} \\")
            painel_anterior = r["panel"]
        linhas.append(f"{r['item']} & {tex(r['label'])} & {r['mean']:.2f} & {r['median']:g} & "
                      f"{r['ratings_1_2_pct']:.1f} & {r['ratings_3_pct']:.1f} & {r['ratings_4_5_pct']:.1f} " + r"\\")
    linhas += [r"\bottomrule", r"\end{tabular}", r"\end{table*}"]
    (saida / "meega_itens.tex").write_text("\n".join(linhas) + "\n", encoding="utf-8")


def argumentos(descricao):
    parser = argparse.ArgumentParser(description=descricao)
    parser.add_argument("--entrada", type=Path, default=ENTRADA, help="Planilha XLSX original.")
    parser.add_argument("--mapa", type=Path, default=BASE / "itens_meega.json", help="Cabeçalhos e rótulos dos itens.")
    parser.add_argument("--saida", type=Path, default=BASE / "resultados", help="Pasta de resultados.")
    return parser.parse_args()


if __name__ == "__main__":
    args = argumentos("Extrair respostas e calcular tabelas do MEEGA+ a partir do Excel.")
    itens, respostas = ler_planilha(args.entrada, args.mapa)
    salvar_tabelas(args.saida, itens, respostas, resumir(itens, respostas))
    print(f"Concluído: {len({r['participant_id'] for r in respostas})} participantes; "
          f"{len(itens)} itens; {len(respostas)} respostas. Saída: {args.saida.resolve()}")
