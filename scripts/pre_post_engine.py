"""Análise descritiva reproduzível do pré/pós-teste de TotalMeltdown.

Uso: python analisar.py --entrada "caminho/respostas.xlsx"
O arquivo de origem é aberto somente para leitura. Veja README.md.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import datetime, timezone
import hashlib
from importlib.metadata import version, PackageNotFoundError
import json
import math
import os
from pathlib import Path
import platform
import statistics as st
import sys
import unicodedata

import openpyxl

HERE = Path(__file__).resolve().parent
DEFAULT_INPUT = "TotalMeltdown_ Questionário de Avaliação (JSS) (respostas).xlsx"
GROUPS = ("none_or_basic", "intermediate_or_advanced")
GROUP_LABELS = {"all": "All participants", "none_or_basic": "None/basic",
                "intermediate_or_advanced": "Intermediate/advanced"}
TRAJECTORIES = ("increased", "unchanged_below_max", "unchanged_at_max", "decreased")


class DataError(ValueError):
    """Entrada incompatível com o protocolo: não gerar resultados parciais."""


def normalize(value):
    # Não remove acentos, não aproxima palavras e não infere equivalência semântica.
    return " ".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()


def check(condition, message):
    if not condition:
        raise DataError(message)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_key(path):
    config = json.loads(Path(path).read_text(encoding="utf-8"))
    items = config["items"]
    check(len(items) == 14, "O protocolo deve conter 14 pares de questões.")
    check([i["pair"] for i in items] == list(range(1, 15)), "Pares do gabarito fora de ordem.")
    for item in items:
        for form in ("pre", "post"):
            q = item[form]
            accepted = {normalize(a) for a in q["accepted_export_answers"]}
            check(normalize(q["correct_export_answer"]) in accepted,
                  f"Gabarito fora do vocabulário aceito: {form}, par {item['pair']}.")
    return config


def resolve_columns(headers, config):
    # Cabeçalhos, não posições fixas: continua correto se as colunas forem reordenadas.
    locations = {}
    for index, header in enumerate(headers):
        if header is not None:
            locations.setdefault(normalize(header), []).append(index)
    required = dict(config["metadata_headers"])
    for item in config["items"]:
        for form in ("pre", "post"):
            required[f"{form}_{item['pair']}"] = item[form]["header"]
    resolved = {}
    for name, header in required.items():
        matches = locations.get(normalize(header), [])
        check(len(matches) == 1,
              f"Cabeçalho ausente ou ambíguo ({name}): {header!r}. Encontrados: {len(matches)}.")
        resolved[name] = matches[0]
    return resolved


def participant_id(value):
    check(value is not None and str(value).strip(), "ID ausente.")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        check(math.isfinite(value) and int(value) == value, f"ID numérico inválido: {value!r}.")
        return str(int(value))
    return str(value).strip()


def score_rows(headers, rows, config):
    columns = resolve_columns(headers, config)
    knowledge = {normalize(k): v for k, v in config["knowledge_groups"].items()}
    consent = normalize(config["accepted_consent"])
    records, scored, seen = [], [], set()
    blank_rows = 0
    maximum = len(config["items"])
    for row_number, row in enumerate(rows, start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            blank_rows += 1
            continue
        def cell(name):
            i = columns[name]
            return row[i] if i < len(row) else None
        pid = participant_id(cell("id"))
        check(pid not in seen, f"ID duplicado na linha {row_number}: {pid}.")
        seen.add(pid)
        check(normalize(cell("consent")) == consent,
              f"Consentimento não confirmado na linha {row_number}; análise interrompida.")
        raw_knowledge = cell("knowledge")
        group = knowledge.get(normalize(raw_knowledge))
        check(group in GROUPS, f"Conhecimento prévio desconhecido na linha {row_number}: {raw_knowledge!r}.")
        rec = dict(participant_id=pid, source_row=row_number,
                   prior_knowledge=str(raw_knowledge).strip(), knowledge_group=group,
                   pre_items=[], post_items=[])
        for item in config["items"]:
            j = item["pair"]
            for form in ("pre", "post"):
                q = item[form]
                response = cell(f"{form}_{j}")
                address = f"{openpyxl.utils.get_column_letter(columns[f'{form}_{j}']+1)}{row_number}"
                check(response is not None and str(response).strip(), f"Resposta ausente em {address} (ID {pid}).")
                answer = normalize(response)
                check(answer in {normalize(a) for a in q["accepted_export_answers"]},
                      f"Resposta não reconhecida em {address} (ID {pid}): {response!r}. Confira o gabarito.json.")
                correct = int(answer == normalize(q["correct_export_answer"]))
                rec[f"{form}_items"].append(correct)
                scored.append(dict(participant_id=pid, knowledge_group=group, form=form,
                                   item=("A" if form=="pre" else "B")+str(j), pair=j,
                                   concept=item["concept_en"], source_cell=address,
                                   response=response, correct_answer=q["correct_export_answer"], score=correct))
        rec["pre_score"] = sum(rec["pre_items"])
        rec["post_score"] = sum(rec["post_items"])
        rec["change"] = rec["post_score"] - rec["pre_score"]
        d = rec["change"]
        rec["trajectory"] = ("increased" if d>0 else "decreased" if d<0 else
                             "unchanged_at_max" if rec["pre_score"]==maximum else "unchanged_below_max")
        rec["pre_at_max"] = int(rec["pre_score"]==maximum)
        rec["changed_item_pairs"] = sum(a!=b for a,b in zip(rec["pre_items"],rec["post_items"]))
        rec["same_total_changed_pattern"] = int(d==0 and rec["changed_item_pairs"]>0)
        records.append(rec)
    check(records, "Nenhum participante encontrado.")
    # Uma linha contém as duas formas do mesmo participante: não parear por posição entre arquivos.
    records.sort(key=lambda r: (0,int(r["participant_id"])) if r["participant_id"].isdigit()
                 else (1,r["participant_id"]))
    return records, scored, columns, blank_rows


def describe(records, group, maximum):
    n = len(records)
    result = dict(group=group, n=n)
    for name in ("pre_score", "post_score", "change"):
        values = [r[name] for r in records]
        result.update({name+"_mean": st.mean(values) if n else None,
                       name+"_median": st.median(values) if n else None,
                       name+"_sd": st.stdev(values) if n>1 else None,
                       name+"_min": min(values) if n else None,
                       name+"_max": max(values) if n else None})
    for state in TRAJECTORIES:
        count = sum(r["trajectory"]==state for r in records)
        result[state+"_n"] = count
        result[state+"_pct"] = 100*count/n if n else None
    result["unchanged_n"] = result["unchanged_below_max_n"]+result["unchanged_at_max_n"]
    result["pre_at_max_n"] = sum(r["pre_score"]==maximum for r in records)
    result["post_at_max_n"] = sum(r["post_score"]==maximum for r in records)
    result["same_total_changed_pattern_n"] = sum(r["same_total_changed_pattern"] for r in records)
    return result


def analyze(records, config):
    maximum = len(config["items"])
    overall = describe(records, "all", maximum)
    groups = [describe([r for r in records if r["knowledge_group"]==g],g,maximum) for g in GROUPS]
    items = []
    for j, item in enumerate(config["items"]):
        transitions = Counter((r["pre_items"][j],r["post_items"][j]) for r in records)
        a = sum(r["pre_items"][j] for r in records)
        b = sum(r["post_items"][j] for r in records)
        n = len(records)
        items.append(dict(pair=j+1, concept=item["concept_en"], domain=item["domain"], n=n,
                          pre_correct=a, post_correct=b, pre_pct=100*a/n, post_pct=100*b/n,
                          difference_n=b-a, difference_pp=100*(b-a)/n,
                          wrong_wrong=transitions[(0,0)], wrong_right=transitions[(0,1)],
                          right_wrong=transitions[(1,0)], right_right=transitions[(1,1)]))
    frequencies = Counter((r["pre_score"],r["post_score"]) for r in records)
    matrix = [dict(pre_score=a,post_score=b,n=frequencies[(a,b)])
              for a in range(maximum+1) for b in range(maximum+1)]
    distribution = []
    for g in ("all",)+GROUPS:
        subset = records if g=="all" else [r for r in records if r["knowledge_group"]==g]
        for score in range(maximum+1):
            distribution.append(dict(group=g,score=score,
                                     pre_n=sum(r["pre_score"]==score for r in subset),
                                     post_n=sum(r["post_score"]==score for r in subset)))
    domains = []
    for domain in dict.fromkeys(i["domain"] for i in config["items"]):
        subset = [i for i in items if i["domain"]==domain]
        total = len(records)*len(subset)
        a,b = sum(i["pre_correct"] for i in subset),sum(i["post_correct"] for i in subset)
        domains.append(dict(domain=domain,items_n=len(subset),participants_n=len(records),
                            response_pairs_n=total,pre_correct=a,post_correct=b,
                            pre_pct=100*a/total,post_pct=100*b/total))
    checks = {
        "participants_equal_group_total": sum(g["n"] for g in groups)==len(records),
        "trajectory_total": sum(overall[s+"_n"] for s in TRAJECTORIES)==len(records),
        "item_transition_totals": all(sum(i[k] for k in ("wrong_wrong","wrong_right","right_wrong","right_right"))==len(records) for i in items),
        "pre_totals": sum(r["pre_score"] for r in records)==sum(i["pre_correct"] for i in items),
        "post_totals": sum(r["post_score"] for r in records)==sum(i["post_correct"] for i in items),
        "paired_change_identity": sum(r["change"] for r in records)==sum(i["wrong_right"]-i["right_wrong"] for i in items),
        "score_matrix_total": sum(v["n"] for v in matrix)==len(records),
    }
    check(all(checks.values()), f"Falha de reconciliação: {checks}")
    return dict(overall=overall,groups=groups,items=items,matrix=matrix,distribution=distribution,
                domains=domains,checks=checks)


def write_csv(path, rows):
    rows = list(rows)
    check(bool(rows), f"Tabela vazia: {path.name}")
    with path.open("w",encoding="utf-8-sig",newline="") as f:
        writer = csv.DictWriter(f,fieldnames=list(rows[0]),delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def tex_escape(text):
    return str(text).replace("&",r"\&").replace("%",r"\%").replace("_",r"\_")


def write_tables(output, records, scored, result, config):
    tables = output/"tabelas"
    tables.mkdir(parents=True,exist_ok=True)
    participants = []
    for r in records:
        row = {k:v for k,v in r.items() if k not in ("pre_items","post_items")}
        row.update({f"A{j+1}":v for j,v in enumerate(r["pre_items"])})
        row.update({f"B{j+1}":v for j,v in enumerate(r["post_items"])})
        participants.append(row)
    exports = {"participantes.csv":participants,"respostas_pontuadas.csv":scored,
               "resumo_geral.csv":[result["overall"]],"grupos_conhecimento.csv":result["groups"],
               "itens.csv":result["items"],"transicoes_escores.csv":result["matrix"],
               "distribuicao_escores.csv":result["distribution"],"blocos_conteudo.csv":result["domains"]}
    applied = []
    for item in config["items"]:
        for form in ("pre","post"):
            q = item[form]
            applied.append(dict(form=form,pair=item["pair"],concept=item["concept_en"],
                                header=q["header"],official_correct_letter=q["official_correct_letter"],
                                official_correct_answer=q["official_correct_answer"],
                                correct_export_answer=q["correct_export_answer"]))
    exports["gabarito_aplicado.csv"] = applied
    for name,rows in exports.items():
        write_csv(tables/name,rows)
    group_tex = [r"\begin{table*}[htbp]",r"\centering",r"\small",
                 r"\caption{Paired test results by self-reported prior knowledge. Scores range from 0 to 14; trajectory columns report participant counts.}",
                 r"\label{tab:knowledge-groups}",r"\begin{tabular}{lrrrrrr}",r"\toprule",
                 r"Prior knowledge & $n$ & Pre mean & Post mean & Increased & Unchanged & Decreased \\",r"\midrule"]
    for g in result["groups"]:
        pre = f"{g['pre_score_mean']:.2f}" if g["n"] else "--"
        post = f"{g['post_score_mean']:.2f}" if g["n"] else "--"
        group_tex.append(f"{GROUP_LABELS[g['group']]} & {g['n']} & {pre} & {post} & {g['increased_n']} & {g['unchanged_n']} & {g['decreased_n']} " + r"\\")
    group_tex += [r"\bottomrule",r"\end{tabular}",r"\end{table*}"]
    (tables/"grupos_conhecimento.tex").write_text("\n".join(group_tex)+"\n",encoding="utf-8")
    item_tex = [r"\begin{table*}[htbp]",r"\centering",r"\small",
                f"\\caption{{Correct answers and transitions across corresponding items in Forms A and B ($n={len(records)}$).}}",
                r"\label{tab:item-transitions}",r"\begin{tabular}{llrrrr}",r"\toprule",
                r"Pair & Concept & Pre & Post & Wrong--right & Right--wrong \\",r"\midrule"]
    for i in result["items"]:
        item_tex.append(f"A{i['pair']}/B{i['pair']} & {tex_escape(i['concept'])} & {i['pre_correct']} & {i['post_correct']} & {i['wrong_right']} & {i['right_wrong']} " + r"\\")
    item_tex += [r"\bottomrule",r"\end{tabular}",r"\end{table*}"]
    (tables/"itens.tex").write_text("\n".join(item_tex)+"\n",encoding="utf-8")


def plot_results(output, result, config):
    os.environ.setdefault("MPLCONFIGDIR",str(output/".matplotlib"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    plt.rcParams.update({"font.family":"DejaVu Sans","font.size":10,"axes.spines.top":False,
                         "axes.spines.right":False,"svg.hashsalt":"totalmeltdown-descriptive-v1",
                         "pdf.fonttype":42,"ps.fonttype":42})
    figures = output/"graficos"
    figures.mkdir(parents=True,exist_ok=True)
    def save(fig,name):
        fig.savefig(figures/(name+".png"),dpi=300,metadata={"Software":"TotalMeltdown analysis"})
        fig.savefig(figures/(name+".svg"),metadata={"Date":None})
        plt.close(fig)
    blue,orange,green,gray = "#2166AC","#C56B16","#27836C","#D7DCE1"
    n = result["overall"]["n"]
    items = result["items"]
    fig,ax = plt.subplots(figsize=(10,7.5))
    fig.subplots_adjust(left=.37,right=.83,top=.83,bottom=.13)
    for y,i in enumerate(items):
        ax.plot([i["pre_pct"],i["post_pct"]],[y,y],color="#A5ABB1",lw=1.5,zorder=1)
    ax.scatter([i["pre_pct"] for i in items],range(14),s=70,facecolors="white",edgecolors=blue,linewidths=1.5,label="Pre-test (Form A)",zorder=3)
    ax.scatter([i["post_pct"] for i in items],range(14),s=32,color=orange,marker="D",label="Post-test (Form B)",zorder=4)
    ax.set_yticks(range(14),[f"A{i['pair']}/B{i['pair']}  {i['concept']}" for i in items])
    ax.set_ylim(13.7,-1.0)
    # Eixo em 0--100% para mostrar também a concentração próxima do teto.
    ax.set_xlim(-2,102)
    ax.set_xticks(range(0,101,20))
    ax.set_xlabel("Correct answers (%)")
    ax.grid(axis="x",color="#E4E7EB",lw=.7)
    ax.set_axisbelow(True)
    ax.tick_params(axis="y",length=0,pad=12)
    ax.spines["left"].set_visible(False)
    for x,label in [(1.07,"Pre"),(1.21,"Post")]:
        ax.text(x,-.7,label,transform=ax.get_yaxis_transform(),ha="center",weight="bold",clip_on=False)
    for y,i in enumerate(items):
        ax.text(1.07,y,str(i["pre_correct"]),transform=ax.get_yaxis_transform(),ha="center",va="center")
        ax.text(1.21,y,str(i["post_correct"]),transform=ax.get_yaxis_transform(),ha="center",va="center")
    fig.suptitle("Correct answers by concept",x=.05,y=.975,ha="left",fontsize=16,weight="bold")
    fig.text(.05,.93,f"Paired responses from {n} students | 14 corresponding item pairs",color="#46515C")
    ax.legend(loc="lower left",bbox_to_anchor=(0,1.04),frameon=False,ncol=2,fontsize=9)
    fig.text(.05,.03,"Forms A and B use different scenarios. Right-hand columns show counts of correct answers.",fontsize=9,color="#46515C")
    save(fig,"01_acertos_por_item")

    summaries = [result["overall"]]+result["groups"]
    labels = [f"{GROUP_LABELS[g['group']]} (n={g['n']})" for g in summaries]
    colors = [green,gray,blue,orange]
    state_labels = ["Increased","Unchanged below 14","Unchanged at 14","Decreased"]
    fig,ax = plt.subplots(figsize=(10,4.8))
    fig.subplots_adjust(left=.30,right=.96,top=.72,bottom=.29)
    for y,g in enumerate(summaries):
        left = 0
        for state,color in zip(TRAJECTORIES,colors):
            count = g[state+"_n"]
            pct = g[state+"_pct"] or 0
            ax.barh(y,pct,left=left,height=.58,color=color,edgecolor="white",lw=1)
            if count:
                ax.text(left+pct/2,y,f"{pct:.1f}%\n({count})",ha="center",va="center",
                        fontsize=9,color="#222222" if color==gray else "white")
            left += pct
    ax.set_yticks(range(len(labels)),labels)
    ax.invert_yaxis()
    ax.set_xlim(0,100)
    ax.set_xlabel("Participants within each group (%)")
    ax.tick_params(axis="y",length=0,pad=10)
    ax.spines["left"].set_visible(False)
    fig.suptitle("Individual score changes by prior knowledge",x=.04,y=.97,ha="left",fontsize=15,weight="bold")
    fig.text(.04,.89,"Prior knowledge is self-reported. Labels show percentages and participant counts.",color="#46515C",fontsize=10)
    fig.legend([Patch(facecolor=c) for c in colors],state_labels,loc="lower center",ncol=2,frameon=False,bbox_to_anchor=(.54,.03),fontsize=10)
    save(fig,"02_trajetorias_por_grupo")

    fig,ax = plt.subplots(figsize=(7.5,6.8))
    fig.subplots_adjust(left=.13,right=.95,top=.80,bottom=.22)
    maximum = len(config["items"])
    ax.plot([0,maximum],[0,maximum],color="#A5ABB1",ls="--",lw=1,zorder=1)
    nonzero = [v for v in result["matrix"] if v["n"]]
    for v in nonzero:
        a,b,count = v["pre_score"],v["post_score"],v["n"]
        color = green if b>a else orange if b<a else blue
        ax.scatter(a,b,s=130+65*count,color=color,edgecolors="white",linewidths=.8,zorder=3)
        ax.text(a,b,str(count),color="white",ha="center",va="center",fontsize=10,weight="bold",zorder=4)
    lower = min(min(v["pre_score"],v["post_score"]) for v in nonzero)
    ax.set_xlim(lower-.6,maximum+.6)
    ax.set_ylim(lower-.6,maximum+.6)
    ax.set_xticks(range(lower,maximum+1))
    ax.set_yticks(range(lower,maximum+1))
    ax.set_xlabel(f"Pre-test score (out of {maximum})")
    ax.set_ylabel(f"Post-test score (out of {maximum})")
    ax.grid(color="#E4E7EB",lw=.7,zorder=0)
    ax.set_axisbelow(True)
    fig.suptitle("Paired score distribution",x=.05,y=.97,ha="left",fontsize=16,weight="bold")
    fig.text(.05,.91,f"n = {n} | Each circle shows the number of students with that score pair.",fontsize=10,color="#46515C")
    fig.text(.05,.055,"The diagonal marks equal scores. Axes display the observed score range.",fontsize=9,color="#46515C")
    fig.legend([Line2D([],[],marker="o",ls="",color=c) for c in (green,blue,orange)],
               ["Increased","Unchanged","Decreased"],loc="lower center",ncol=3,frameon=False,bbox_to_anchor=(.53,.095))
    save(fig,"03_distribuicao_pareada")


def write_summary(output, result, config):
    a = result["overall"]
    lines = ["# Resultados descritivos do pré e pós-teste", "",
             f"Participantes: {a['n']}. Máximo: {len(config['items'])} pontos por teste.","",
             f"Média no pré: {a['pre_score_mean']:.4f}; no pós: {a['post_score_mean']:.4f}; variação média: {a['change_mean']:.4f} ponto.",
             f"Aumentaram: {a['increased_n']}; mantiveram: {a['unchanged_n']}; diminuíram: {a['decreased_n']}.",
             f"Teto no pré: {a['pre_at_max_n']}; teto no pós: {a['post_at_max_n']}; máximo nos dois: {a['unchanged_at_max_n']}.",
             f"Mesma nota total com mudança no padrão de acertos/erros: {a['same_total_changed_pattern_n']}.","",
             "| Conhecimento autodeclarado | n | Média pré | Média pós | Aumentaram | Mantiveram | Diminuíram |",
             "|---|---:|---:|---:|---:|---:|---:|"]
    for g in result["groups"]:
        pre = f"{g['pre_score_mean']:.4f}" if g["n"] else "—"
        post = f"{g['post_score_mean']:.4f}" if g["n"] else "—"
        lines.append(f"| {GROUP_LABELS[g['group']]} | {g['n']} | {pre} | {post} | {g['increased_n']} | {g['unchanged_n']} | {g['decreased_n']} |")
    lines += ["", "A análise é descritiva e exploratória. As formas A e B usam cenários diferentes; os pares representam conteúdos correspondentes e não têm equivalência psicométrica demonstrada. As trajetórias descrevem pontuações, sem estabelecer aprendizagem causal ou retenção. Os grupos representam autodeclaração de conhecimento, não níveis de proficiência comprovados.","",
              "As contagens por questão representam participantes; os blocos de conteúdo agregam respostas a várias questões por pessoa e não aumentam o tamanho da amostra. Nenhum participante foi removido por nota alta, queda de pontuação ou tamanho da variação.","",
              "Não foram calculados testes de significância, correlações pré-versus-variação, ganho normalizado, modelos de regressão ou intervalos bootstrap. Consulte README.md para o protocolo e os dicionários dos arquivos.",""]
    (output/"resumo.md").write_text("\n".join(lines),encoding="utf-8")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--entrada",type=Path,default=HERE.parent/DEFAULT_INPUT)
    parser.add_argument("--gabarito",type=Path,default=HERE/"gabarito.json")
    parser.add_argument("--aba",help="Nome exato da aba; por padrão usa o nome registrado no gabarito.")
    parser.add_argument("--saida",type=Path,default=HERE/"resultados")
    parser.add_argument("--sem-graficos",action="store_true",help="Produzir somente tabelas e resumo, sem importar Matplotlib.")
    args = parser.parse_args(argv)
    try:
        source = args.entrada.resolve()
        config = load_key(args.gabarito)
        original_hash = sha256(source)
        workbook = openpyxl.load_workbook(source,read_only=True,data_only=True)
        try:
            sheet_name = args.aba or config["sheet_name"]
            check(sheet_name in workbook.sheetnames,f"Aba não encontrada: {sheet_name!r}.")
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator)
            records,scored,columns,blank_rows = score_rows(headers,iterator,config)
        finally:
            workbook.close()
        result = analyze(records,config)
        output = args.saida.resolve()
        check(output != source, "A saída deve ser um diretório diferente da planilha.")
        output.mkdir(parents=True,exist_ok=True)
        write_tables(output,records,scored,result,config)
        write_summary(output,result,config)
        if not args.sem_graficos:
            plot_results(output,result,config)
        check(sha256(source)==original_hash,"O arquivo de origem mudou durante a execução.")
        versions = {}
        for package in ("openpyxl","matplotlib","numpy","Pillow","contourpy","cycler","fonttools","kiwisolver","packaging","pyparsing","python-dateutil","six","et-xmlfile"):
            try:
                versions[package] = version(package)
            except PackageNotFoundError:
                pass
        expected_dirs = [output/"tabelas"]+([] if args.sem_graficos else [output/"graficos"])
        output_files = [p for folder in expected_dirs for p in sorted(folder.iterdir()) if p.is_file()]+[output/"resumo.md"]
        metadata = dict(protocol_version=config["protocol_version"], executed_utc=datetime.now(timezone.utc).isoformat(),
                        python=platform.python_version(),platform=platform.platform(),packages=versions,
                        source_file=source.name,source_sha256=original_hash,source_sheet=sheet_name,
                        script_sha256=sha256(__file__),key_sha256=sha256(args.gabarito),
                        participants_n=len(records),scored_responses_n=len(scored),blank_rows_ignored=blank_rows,
                        exclusions=0,graphics_generated=not args.sem_graficos,
                        columns={k:openpyxl.utils.get_column_letter(v+1) for k,v in columns.items()},
                        checks=result["checks"],outputs={str(p.relative_to(output)).replace("\\","/"):sha256(p) for p in output_files})
        (output/"metadados.json").write_text(json.dumps(metadata,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        (output/"resultados.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        a=result["overall"]
        print(f"Concluído: {a['n']} participantes; {len(scored)} respostas pontuadas.")
        print(f"Pré: {a['pre_score_mean']:.4f}; pós: {a['post_score_mean']:.4f}; variação: {a['change_mean']:.4f}.")
        print(f"Aumentaram/mantiveram/diminuíram: {a['increased_n']}/{a['unchanged_n']}/{a['decreased_n']}.")
        print(f"Saída: {output}")
        return 0
    except (DataError,OSError,KeyError,StopIteration,json.JSONDecodeError,ImportError) as error:
        print(f"ERRO: {error}",file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
